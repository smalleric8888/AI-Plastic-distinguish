import os
import openpyxl

BASE_DATA_DIR = r"C:\Users\small\Desktop\AI Agent 待處理\天辰"
EXCEL_PATH = os.path.join(BASE_DATA_DIR, "天辰廢料全循環履歷系統品項(1).xlsx")
SOUND_DIR = os.path.join(BASE_DATA_DIR, "聲音檔")
VIDEO_DIR = os.path.join(BASE_DATA_DIR, "影片檔")
PHOTO_APP_DIR = os.path.join(BASE_DATA_DIR, "生產履歷APP照片115.8.20")
TEST_PHOTO_DIR = os.path.join(BASE_DATA_DIR, "測試照片")

EMBEDDED_RULES = {
    "ABS": {
        "材質": "ABS-丙烯腈-丁二烯-苯乙烯共聚物",
        "市場常用名稱": "ABS 樹脂、超高衝擊塑料、樂高積木塑料",
        "產源產業別": "家電/3C製造業、汽機車零配件業",
        "原始用途": "安全帽、家電外殼、鍵盤、汽車保桿",
        "水中浮沉": "沉於水（密度 ~1.05 g/cm³）",
        "燃燒特徵": "燒焦輪胎味、帶甜感化學橡膠味，火焰呈黃色，濃黑煙",
        "關鍵字": "1. 電子電器外殼  2. 汽車零配件  3. 3D列印線材"
    },
    "PVC": {
        "材質": "聚氯乙烯",
        "市場常用名稱": "PVC",
        "產源產業別": "建材營造業、配管業、人造皮革業",
        "原始用途": "水管、門窗框、雨衣、軟質膠布/地板",
        "水中浮沉": "沉於水（密度 1.30–1.45 g/cm³）",
        "燃燒特徵": "強烈鹽酸酸臭味、焦糊橡膠味，火焰下端綠色，離火自熄",
        "關鍵字": "建築水管、電線外皮、雨衣、醫療導管"
    },
    "PS": {
        "材質": "聚苯乙烯",
        "市場常用名稱": "6號塑膠、保麗龍、硬膠",
        "產源產業別": "免洗餐具業、包裝緩衝材、電子外殼",
        "原始用途": "免洗餐具、養樂多瓶、緩衝保麗龍、CD盒",
        "水中浮沉": "沉於水（密度 ~1.05 g/cm³）",
        "燃燒特徵": "怪異甜香轉焦苦橡膠臭，亮橘黃火焰，極濃黑煙與炭黑飛揚",
        "關鍵字": "免洗餐具、養樂多瓶、保麗龍緩衝材、CD盒"
    },
    "PP": {
        "材質": "聚丙烯",
        "市場常用名稱": "5號塑膠",
        "產源產業別": "食品容器包裝業、紡織編織業、汽配業",
        "原始用途": "飲料杯、保鮮盒、編織袋、汽車內裝",
        "水中浮沉": "浮於水（密度 ~0.90–0.91 g/cm³）",
        "燃燒特徵": "剛點燃有石蠟味，持續燃燒轉柴油/機油味，白煙極少",
        "關鍵字": "微波保鮮盒、飲料杯蓋、汽車保險桿、收納箱"
    },
    "PE": {
        "材質": "聚乙烯",
        "市場常用名稱": "塑膠骨、塑膠粒、HDPE、LDPE",
        "產源產業別": "物流包裝業、日化容器吹瓶業、農膜業",
        "原始用途": "鮮奶瓶、清潔劑桶、塑膠袋、氣泡膜",
        "水中浮沉": "浮於水（密度 0.91–0.96 g/cm³）",
        "燃燒特徵": "強烈燒焦蠟燭/滴蠟石蠟味，藍底黃尖火焰，幾乎無煙",
        "關鍵字": "塑膠袋、保鮮膜、鮮奶瓶、大貨桶、溫室農膜"
    }
}

class PlasticDataLoader:
    def __init__(self, excel_path=EXCEL_PATH):
        self.excel_path = excel_path
        self.materials_data = {}
        self.sample_items = []
        self.photo_mapping = {}
        self.sound_files = {}
        self.video_files = {}
        self._load_data()

    def _load_data(self):
        # 1. 載入 Excel (或內建雲端備份庫)
        loaded_excel = False
        if os.path.exists(self.excel_path):
            try:
                wb = openpyxl.load_workbook(self.excel_path)
                if "辨別方式總表" in wb.sheetnames:
                    ws = wb["辨別方式總表"]
                    current_mat = None
                    for r in range(2, ws.max_row + 1):
                        mat_code = ws.cell(r, 2).value
                        cat = ws.cell(r, 3).value
                        fld = ws.cell(r, 4).value
                        val = ws.cell(r, 5).value

                        if mat_code:
                            current_mat = str(mat_code).strip()
                            if current_mat not in self.materials_data:
                                self.materials_data[current_mat] = {}

                        if current_mat and fld:
                            field_name = str(fld).strip()
                            field_val = str(val).strip() if val is not None else ""
                            self.materials_data[current_mat][field_name] = field_val
                    loaded_excel = True
            except Exception:
                pass

        if not loaded_excel or len(self.materials_data) == 0:
            self.materials_data = EMBEDDED_RULES.copy()

        # 2. 掃描聲音檔
        if os.path.exists(SOUND_DIR):
            for f in os.listdir(SOUND_DIR):
                full_path = os.path.join(SOUND_DIR, f)
                if os.path.isfile(full_path):
                    name_lower = f.lower()
                    for mat in ["ABS", "PE", "PP", "PVC", "PET", "PS"]:
                        if mat.lower() in name_lower:
                            if mat not in self.sound_files:
                                self.sound_files[mat] = []
                            self.sound_files[mat].append({"filename": f, "path": full_path})
        else:
            # 雲端版預設媒體標記 (共 13 個聲音標記)
            self.sound_files = {
                "ABS": [{"filename": "ABS 管敲擊聲.m4a", "path": ""}],
                "PE": [{"filename": "PE 塞頭本色.m4a", "path": ""}, {"filename": "PE 大白.m4a", "path": ""}, {"filename": "PE 中白低密度.m4a", "path": ""}, {"filename": "PE雜色（藍色）.m4a", "path": ""}, {"filename": "PE雜色（黑色）.m4a", "path": ""}],
                "PP": [{"filename": "PP塞頭本色.m4a", "path": ""}, {"filename": "PP大白.m4a", "path": ""}, {"filename": "PP打包帶.m4a", "path": ""}],
                "PVC": [{"filename": "PVC管敲擊聲.m4a", "path": ""}],
                "PET": [{"filename": "PET 打包帶.m4a", "path": ""}, {"filename": "PET 膜.m4a", "path": ""}, {"filename": "PET2.m4a", "path": ""}]
            }

        # 3. 掃描影片檔
        if os.path.exists(VIDEO_DIR):
            for f in os.listdir(VIDEO_DIR):
                full_path = os.path.join(VIDEO_DIR, f)
                if os.path.isfile(full_path):
                    name_lower = f.lower()
                    for mat in ["ABS", "PE", "PP", "PVC", "PET", "PS"]:
                        if mat.lower() in name_lower:
                            if mat not in self.video_files:
                                self.video_files[mat] = []
                            self.video_files[mat].append({"filename": f, "path": full_path})
        else:
            # 雲端版預設影片標記 (共 8 個影片標記)
            self.video_files = {
                "ABS": [{"filename": "ABS 燃燒.mp4", "path": ""}],
                "PE": [{"filename": "PE 塞頭本色燃燒.mp4", "path": ""}, {"filename": "PE 大白燃燒.mp4", "path": ""}],
                "PP": [{"filename": "PP塞頭本色燃燒.mp4", "path": ""}, {"filename": "PP打包帶燃燒.mp4", "path": ""}],
                "PVC": [{"filename": "PVC管燃燒.mp4", "path": ""}],
                "PET": [{"filename": "PET打包帶燃燒.mp4", "path": ""}, {"filename": "PET2 燃燒.mp4", "path": ""}]
            }

        # 4. 掃描樣品庫
        if os.path.exists(PHOTO_APP_DIR):
            for folder in os.listdir(PHOTO_APP_DIR):
                folder_path = os.path.join(PHOTO_APP_DIR, folder)
                if os.path.isdir(folder_path):
                    imgs = [os.path.join(folder_path, f) for f in os.listdir(folder_path) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
                    auds = [os.path.join(folder_path, f) for f in os.listdir(folder_path) if f.lower().endswith(('.m4a', '.mp3', '.wav'))]
                    vids = [os.path.join(folder_path, f) for f in os.listdir(folder_path) if f.lower().endswith(('.mp4', '.avi', '.mov'))]
                    
                    mat_key = "其他"
                    for m in ["ABS", "PET", "PE", "PP", "PVC", "PS"]:
                        if m.lower() in folder.lower():
                            mat_key = m
                            break

                    self.sample_items.append({
                        "name": folder,
                        "material": mat_key,
                        "folder_path": folder_path,
                        "images": imgs,
                        "audios": auds,
                        "videos": vids
                    })
        else:
            # 雲端版預設對照樣品組 (共 14 組樣品)
            sample_names = [
                "ABS管", "PET2", "PET打包帶", "PET膜", 
                "PE中白(低密度)", "PE塞頭本色", "PE大白", "PE膜雜色(藍色)", "PE膜雜色(黑色)",
                "PP塞頭本色", "PP塞頭雜色", "PP大白", "PP打包帶", "PVC管"
            ]
            for m in sample_names:
                self.sample_items.append({
                    "name": m,
                    "material": m[:3],
                    "folder_path": "",
                    "images": [],
                    "audios": [],
                    "videos": []
                })

    def get_material_info(self, mat_code):
        return self.materials_data.get(mat_code, EMBEDDED_RULES.get(mat_code, {}))

    def get_reference_media(self, mat_code):
        matched_sounds = self.sound_files.get(mat_code, [])
        matched_videos = self.video_files.get(mat_code, [])
        matched_samples = [s for s in self.sample_items if s["material"] == mat_code]
        return {
            "sounds": matched_sounds,
            "videos": matched_videos,
            "samples": matched_samples
        }

if __name__ == "__main__":
    loader = PlasticDataLoader()
    print("Materials loaded:", list(loader.materials_data.keys()))
